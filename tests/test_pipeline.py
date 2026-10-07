import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, rules
from app.main import app

c = TestClient(app)


def run(query=""):
    return {r["stage"]: r for r in c.post(f"/api/run-all{query}").json()["runs"]}


def controls():
    return {x["id"]: x["status"] for x in c.get("/controls/output").json()["controls"]}


def test_counts_on_sample_data():
    r = run()
    assert r["dao"]["rows_out"] == 32
    y = r["yxy"]
    assert (y["rows_out"], y["filtered"], y["duplicates_removed"], y["rejected"]) == (24, 4, 1, 3)
    assert r["aoc"]["rows_out"] == r["mis"]["rows_out"] == 24


def test_controls_pass_on_clean_run():
    run()
    ctl = controls()
    assert ctl["C1"] == ctl["C2"] == ctl["C3"] == ctl["C4"] == "PASS"


def test_controls_catch_injected_fault():
    run("?inject_fault=true")
    ctl = controls()
    assert ctl["C2"] == ctl["C3"] == ctl["C4"] == "FAIL"


def test_settlement_is_outstanding_plus_interest():
    run()
    for a in c.get("/aoc/output").json()["rows"]:
        assert abs(a["settlement_amount"] - (a["outstanding_bal"] + a["interest_amount"])) < 0.006


def test_pif_rule():
    assert rules.pif_flag(0, 0, 1, 0.0) == 1
    assert rules.pif_flag(1, 0, 1, 0.0) == 0      # still active
    assert rules.pif_flag(0, 1, 1, 0.0) == 0      # charged off
    assert rules.pif_flag(0, 0, 2, 0.0) == 0      # not a paid-in-full closure code
    assert rules.pif_flag(0, 0, 1, 500.0) == 0    # balance remaining


def test_pass_through_values_are_not_altered():
    run()
    dao = {r["clrty_id"]: r for r in c.get("/dao/output").json()["rows"]}
    for m in c.get("/mis/output").json()["rows"]:
        assert m["Treatment"] == dao[m["Clarity ID"]]["workout_completed_cd"]
        assert m["Outstanding Balance"] == dao[m["Clarity ID"]]["settlement_am"]


def test_stage_order_enforced():
    assert c.post("/yxy/run").status_code == 409
    assert c.post("/aoc/run").status_code == 409
    assert c.post("/mis/run").status_code == 409
    assert c.post("/controls/run").status_code == 409


def test_results_persist_in_sqlite():
    run()
    assert len(db.get_rows("mis")) == 24
    assert db.get_kv("runs")["mis"]["rows_out"] == 24


def test_parameters_change_the_filter():
    r = run("?file_month=2026-08-01")
    assert r["yxy"]["rows_out"] == 2 and r["yxy"]["filtered"] == 30


def test_trace_filtered_rejected_and_good():
    run()
    assert c.get("/api/trace/100025").json()["filtered_out"] is True
    assert c.get("/api/trace/100030").json()["rejected"][0]["reason"] == "NEGATIVE_SETTLEMENT_AM"
    assert c.get("/api/trace/100031").json()["rejected"][0]["reason"] == "NO_ACCOUNT_MASTER_ROW"
    assert c.get("/api/trace/100004").json()["reached"] == ["dao", "yxy", "aoc", "mis"]
    assert c.get("/api/trace/999999").status_code == 404


def test_csv_round_trip_gives_same_result():
    before = run()["mis"]["notes"]
    sample = c.get("/dao/sample.csv").content
    acc = c.get("/yxy/account-master.csv").content
    assert c.post("/dao/upload", files={"file": ("d.csv", sample)}).status_code == 200
    assert c.post("/yxy/upload-account-master", files={"file": ("a.csv", acc)}).status_code == 200
    assert c.get("/api/status").json()["sources"]["dao"]["kind"] == "uploaded"
    assert run()["mis"]["notes"] == before
    c.post("/api/reset-sources")
    assert c.get("/api/status").json()["sources"]["dao"]["kind"] == "sample"


def test_bad_csv_is_rejected_with_a_clear_message():
    r = c.post("/dao/upload", files={"file": ("x.csv", b"clrty_id,foo\n1,2\n")})
    assert r.status_code == 422 and "Missing required column" in r.json()["detail"]
    header = c.get("/dao/sample.csv").text.splitlines()[0]
    bad = header + "\nabc" + ",x" * 11 + "\n"
    r = c.post("/dao/upload", files={"file": ("x.csv", bad.encode())})
    assert r.status_code == 422 and "clrty_id" in r.json()["detail"]


def test_excel_export_has_report_and_controls():
    run()
    r = c.get("/mis/download")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["MIS_14MQ", "Controls"]
    assert wb["MIS_14MQ"].max_row == 25      # header + 24 rows


def test_catalog_flags_known_gaps_in_the_sheet():
    run()
    issues = " ".join(f["issue"] for f in c.get("/catalog/output").json()["findings"])
    assert "settment_amount" in issues and "references 'settlement_amount'" in issues and "blank" in issues


@pytest.mark.parametrize("cde", [m["cde"] for m in rules.CDE_MAP])
def test_mermaid_generated_for_every_cde(cde):
    assert c.get(f"/catalog/mermaid/{cde}").json()["mermaid"].startswith("flowchart LR")
