"""App 5 - MIS Report. Builds the MIS_14MQ report with the report data element names and exports it to Excel."""
from io import BytesIO
from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from app import db, rules

app = FastAPI(title="5 - MIS Report", description="Final report layer, with Excel export.")


@app.post("/run")
def run():
    aoc = db.get_rows("aoc")
    if not aoc:
        raise HTTPException(409, "AOC layer is empty. Run AOC first.")
    out = [{mis: r[a] for a, mis in rules.AOC_TO_MIS} for r in aoc]
    db.set_rows("mis", out)
    total = round(sum(r["Settlement Amount"] for r in out), 2)
    return db.set_run("mis", {"stage": "mis", "run_at": db.now(), "rows_in": len(aoc), "rows_out": len(out), "rejected": 0,
                              "notes": [f"Total Settlement Amount {total:,.2f}", f"PIF accounts: {sum(r['PIF'] for r in out)}"]})


@app.get("/output")
def output():
    return {"layer": "MIS Report", "rows": db.get_rows("mis")}


@app.get("/download")
def download():
    rows = db.get_rows("mis")
    if not rows:
        raise HTTPException(409, "MIS report is empty. Run the pipeline first.")
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "MIS_14MQ"
    cols = [m for _, m in rules.AOC_TO_MIS]
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, c in enumerate(cols, 1):
        ws.cell(1, i).font = Font(bold=True)
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(14, len(c) + 4)
    ctl = db.get_kv("controls")
    if ctl:
        w2 = wb.create_sheet("Controls")
        w2.append(["Control", "Check", "Status", "Checked", "Failed"])
        for c in ctl["controls"]:
            w2.append([c["id"], c["name"], c["status"], c["checked"], c["failed"]])
        for c in w2[1]:
            c.font = Font(bold=True)
        w2.column_dimensions["B"].width = 70
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="MIS_14MQ.xlsx"'})
